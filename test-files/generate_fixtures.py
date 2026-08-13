"""Regenerate every binary fixture under test-files/.

These are real files (real PDF bytes, real .xlsx bytes, a real JPEG) --
not synthetic in-memory objects like backend/tests/test_usp_helpers.py
uses -- so they exercise the actual text-extraction and parsing code
paths (app/services/document_text.py, app/services/any_form.py, etc.)
the same way a professor's upload would.

Run from anywhere:
    backend/.venv/Scripts/python.exe test-files/generate_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pymupdf
from docx import Document
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).parent


def write_pdf(path: Path, lines: list[str]) -> None:
    document = pymupdf.open()
    page = document.new_page()
    text = "\n".join(lines)
    page.insert_text((72, 72), text, fontsize=11, lineheight=1.4)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    document.close()


def write_docx(path: Path, lines: list[str]) -> None:
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def write_image_cv(path: Path, lines: list[str]) -> None:
    """A JPEG photo/scan of a CV -- no embedded/selectable text at all, the
    same as a phone photo of a printed page. CV Import reads this via the
    vision-model OCR path (extract_text_with_ocr_fallback), not deterministic
    extraction -- there's no text layer to read any other way. A large
    default font is used because OCR accuracy on a synthetic screenshot-style
    render is otherwise unrepresentatively worse than on an actual photo."""

    try:
        font = ImageFont.truetype("arial.ttf", 28)
    except OSError:
        font = ImageFont.load_default()
    image = Image.new("RGB", (1000, 900), color="white")
    draw = ImageDraw.Draw(image)
    y = 40
    for line in lines:
        draw.text((40, y), line, fill="black", font=font)
        y += 44
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, format="JPEG", quality=92)


def write_scanned_pdf(path: Path, lines: list[str]) -> None:
    """A PDF with NO text layer at all -- an image of text embedded as a page
    background, exactly what "scan to PDF" apps and old flatbed scanners
    produce. pymupdf's get_text() returns empty on this, which is precisely
    the case extract_text_with_ocr_fallback's page-image-OCR path exists for.
    """

    try:
        font = ImageFont.truetype("arial.ttf", 28)
    except OSError:
        font = ImageFont.load_default()
    image = Image.new("RGB", (1000, 1300), color="white")
    draw = ImageDraw.Draw(image)
    y = 40
    for line in lines:
        draw.text((40, y), line, fill="black", font=font)
        y += 44
    image_path = path.with_suffix(".page.jpg")
    image.save(image_path, format="JPEG", quality=92)

    document = pymupdf.open()
    page = document.new_page(width=1000, height=1300)
    page.insert_image(pymupdf.Rect(0, 0, 1000, 1300), filename=str(image_path))
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    document.close()
    image_path.unlink()


def write_xlsx(path: Path, rows: list[list[object]]) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


CV_LINES = [
    "Dr. Ananya Sharma",
    "Associate Professor, Department of Computer Science",
    "",
    "ACADEMIC ACTIVITIES",
    "Attended FDP on Generative AI, IIT Bombay, 2024",
    "Delivered an invited talk on Federated Learning at IEEE Pune Chapter, 2023",
    "Published a paper titled 'Privacy-Preserving ML' in IEEE Transactions, 2023",
    "Reviewer for the International Conference on Machine Learning, 2022",
    "Supervised 3 undergraduate capstone projects, 2021-2024",
    "Member, Departmental Curriculum Committee, 2020-present",
    "Received the Best Paper Award at NCVPRIPG, 2019",
    "",
    "This line has no academic keyword and should not become an activity.",
]


def main() -> None:
    write_pdf(ROOT / "cv-import" / "sample_cv.pdf", CV_LINES)
    write_docx(ROOT / "cv-import" / "sample_cv.docx", CV_LINES)
    write_xlsx(ROOT / "cv-import" / "sample_cv.xlsx", [[line] for line in CV_LINES])
    write_image_cv(ROOT / "cv-import" / "photographed_cv.jpg", CV_LINES)
    write_scanned_pdf(ROOT / "cv-import" / "scanned_cv.pdf", CV_LINES)
    # Still a genuine edge case: a plain-text export is not in
    # ALLOWED_CV_MIME_TYPES, so this must be rejected at upload (415), not
    # silently accepted and then fail deeper in the pipeline.
    (ROOT / "cv-import" / "unsupported_cv.txt").parent.mkdir(parents=True, exist_ok=True)
    (ROOT / "cv-import" / "unsupported_cv.txt").write_text("\n".join(CV_LINES), encoding="utf-8")

    write_xlsx(
        ROOT / "any-form" / "matching_schema.xlsx",
        [
            ["Annual Faculty Appraisal Form"],
            [],
            ["Faculty Name:", ""],
            ["Employee Code:", ""],
            ["Department:", ""],
            ["Email:", ""],
            ["Number of Publications:", ""],
            ["Number of FDPs:", ""],
            ["Committees:", ""],
            ["Random note (not a field, no trailing colon)", "ignore me"],
        ],
    )
    # Edge case: no "Label:" cells at all -- 0 fields should be detected,
    # not an error. A real professor could easily upload the wrong file.
    write_xlsx(
        ROOT / "any-form" / "edge_case_no_labels.xlsx",
        [
            ["Month", "Hours Taught"],
            ["August", 40],
            ["September", 38],
        ],
    )

    write_pdf(ROOT / "teaching-change" / "course_v1" / "syllabus.pdf", [
        "CS301 Syllabus", "Week 1: Introduction", "Week 2: Basics of Algorithms",
    ])
    write_pdf(ROOT / "teaching-change" / "course_v1" / "lab1.pdf", ["Lab 1: Sorting exercise"])
    write_pdf(ROOT / "teaching-change" / "course_v2" / "syllabus.pdf", [
        "CS301 Syllabus", "Week 1: Introduction", "Week 2: Basics of Algorithms",
        "Week 3: Advanced Topics in Deep Learning",
    ])
    write_pdf(ROOT / "teaching-change" / "course_v2" / "lab1.pdf", ["Lab 1: Sorting exercise"])
    write_pdf(ROOT / "teaching-change" / "course_v2" / "lab2_new.pdf", ["Lab 2: Cloud deployment exercise"])

    write_pdf(ROOT / "evidence-library" / "sample_certificate.pdf", [
        "Certificate of Appreciation", "Presented to Dr. Ananya Sharma",
        "For the Invited Talk at IEEE Pune Chapter", "14 October 2025",
    ])
    write_image_cv(ROOT / "evidence-library" / "sample_certificate.jpg", [
        "Certificate of Appreciation", "Dr. Ananya Sharma", "IEEE Pune Chapter",
    ])

    print("Fixtures written under", ROOT)


if __name__ == "__main__":
    main()
